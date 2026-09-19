from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageOps


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "project-documentation"
SCREEN_DIR = OUT_DIR / "screenshots"
DOCX_IMG_DIR = SCREEN_DIR / "docx"
DOCX_PATH = OUT_DIR / "Agentic_AI_Code_Review_Project_Documentation.docx"


def main():
    DOCX_IMG_DIR.mkdir(parents=True, exist_ok=True)
    prepared_images = prepare_images()

    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_document_control(doc)
    add_executive_summary(doc)
    add_architecture(doc)
    add_workflow(doc)
    add_screenshots(doc, prepared_images)
    add_components(doc)
    add_api_and_models(doc)
    add_security_observability(doc)
    add_operations(doc)
    add_ground_truth(doc)
    add_testing_and_acceptance(doc)
    add_limitations_and_roadmap(doc)

    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size in [("Title", 24), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)]:
        style = styles[style_name]
        style.font.name = "Aptos"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Agentic AI Code Review and Quality Engineering Platform")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 99, 117)


def add_cover(doc):
    title = doc.add_paragraph(style="Title")
    title.add_run("Agentic AI Code Review and Quality Engineering Platform Project Documentation")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run("Technical architecture, workflow, operations, security, quality checks, and demonstration guide")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(71, 84, 103)

    p = doc.add_paragraph()
    p.add_run("Purpose. ").bold = True
    p.add_run(
        "This document explains the working TypeScript, NestJS, and React demonstration project. "
        "It is written for technical architects, senior engineers, and engineering leaders who need to understand how the platform combines agentic AI review with deterministic quality engineering."
    )

    add_key_value_table(
        doc,
        [
            ("Project type", "Working local demonstration platform"),
            ("Primary stack", "TypeScript, NestJS, React, Vite, Node.js"),
            ("AI provider", "Gemini API with local structured fallback"),
            ("Quality tools", "ESLint, TypeScript compiler, Jest coverage, npm audit"),
            ("Current policy outcome for sample project", "BLOCKED because deterministic gates fail"),
        ],
    )


def add_document_control(doc):
    doc.add_heading("Document Control", level=1)
    add_table(
        doc,
        ["Item", "Value"],
        [
            ["Document owner", "Engineering architecture demo"],
            ["Audience", "Technical architects, senior engineers, engineering leaders, reviewers"],
            ["Repository", "agentic-review-demo"],
            ["Primary deliverables", "React UI, NestJS API, agents package, shared schemas, sample project, docs"],
            ["Review status", "Platform build, tests, lint, and end-to-end API flow verified"],
        ],
        widths=[1.8, 4.9],
    )


def add_executive_summary(doc):
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "The platform reviews a local repository through a NestJS orchestrator. It gathers Git context, runs an AI Code Review Agent, runs a deterministic Quality Agent, aggregates findings, evaluates an explicit policy, and displays the outcome in a React dashboard."
    )
    doc.add_paragraph(
        "The core architectural decision is separation of authority. Deterministic tools remain authoritative for compiler, lint, test, coverage, and dependency facts. AI contributes contextual findings with evidence, confidence, file references, and suggested fixes. AI findings require human validation and do not directly block production."
    )
    add_table(
        doc,
        ["Objective", "How the project demonstrates it"],
        [
            ["Agentic AI", "The Code Review Agent uses controlled repository tools and emits structured findings."],
            ["Quality engineering", "The Quality Agent executes real deterministic checks through a command allowlist."],
            ["Enterprise guardrails", "Path validation, schema validation, environment-only secrets, and read-only AI access are implemented."],
            ["Human control", "Reviewers can accept, dismiss, or resolve findings from the dashboard."],
            ["Observability", "Operational agent and tool events are captured as a trace without exposing hidden model reasoning."],
        ],
        widths=[1.9, 4.8],
    )


def add_architecture(doc):
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph(
        "The repository is organized as a compact monorepo. Application code is separated from reusable agent and schema packages, while the intentionally defective sample project remains isolated from platform quality checks."
    )
    add_table(
        doc,
        ["Area", "Path", "Responsibility"],
        [
            ["Frontend", "apps/web", "React dashboard for review execution, findings, quality checks, and trace."],
            ["Backend", "apps/api", "NestJS API, review lifecycle, orchestration, and finding updates."],
            ["Agents", "packages/agents", "Code Review Agent, Quality Agent, repository tools, policy engine, aggregator, trace recorder."],
            ["Shared schemas", "packages/shared", "Zod schemas and TypeScript types shared by API, agents, and UI."],
            ["Sample app", "sample-project", "Small NestJS Orders domain with intentional defects for demonstration."],
            ["Documentation", "docs/project-documentation", "Project document, screenshots, and Word deliverable."],
        ],
        widths=[1.15, 1.55, 4.0],
    )
    doc.add_heading("Logical Flow", level=2)
    add_numbered_list(
        doc,
        [
            "Developer opens the dashboard and starts a review.",
            "NestJS validates the repository path and collects Git context.",
            "The Quality Agent runs deterministic tools through an allowlist.",
            "The Code Review Agent gathers bounded repository evidence and creates structured AI findings.",
            "The aggregator deduplicates findings from AI and deterministic sources.",
            "The policy engine produces PASS, WARNING, HUMAN REVIEW REQUIRED, or BLOCKED.",
            "The dashboard shows the decision, quality checks, findings, and trace.",
            "A human reviewer accepts, dismisses, or resolves findings.",
        ],
    )


def add_workflow(doc):
    doc.add_heading("Review Workflow", level=1)
    doc.add_paragraph(
        "The runtime workflow is intentionally visible in the dashboard so a reviewer can see what the system is doing and why a decision was made."
    )
    add_table(
        doc,
        ["Step", "Runtime activity", "Source of truth"],
        [
            ["1", "Initialize review and create trace recorder", "NestJS orchestrator"],
            ["2", "Validate repository and collect Git context", "Repository tools"],
            ["3", "Run TypeScript, ESLint, Jest, coverage, and npm audit", "Quality Agent"],
            ["4", "Analyze changed code and related files", "Code Review Agent"],
            ["5", "Normalize and deduplicate findings", "Review Aggregator"],
            ["6", "Apply deterministic and AI policy rules", "Policy Engine"],
            ["7", "Show findings, quality checks, and trace", "React dashboard"],
            ["8", "Accept, dismiss, or resolve findings", "Human reviewer"],
        ],
        widths=[0.55, 4.0, 2.15],
    )


def add_screenshots(doc, images):
    doc.add_heading("Running Project Screenshots", level=1)
    doc.add_paragraph(
        "These screenshots were captured from the running application and show the current demo state."
    )
    screenshot_items = [
        ("Dashboard before review", "The initial dashboard shows repository context and the primary review action.", "01-dashboard-ready.png"),
        ("Dashboard after completed review", "The completed review shows the BLOCKED policy decision and live review summary.", "02-dashboard-review-complete.png"),
        ("Quality checks panel", "Deterministic quality outputs are displayed with real status, duration, metrics, and summaries.", "03-quality-checks.png"),
        ("Findings and human review", "AI and deterministic findings are shown with severity, source, evidence, and human actions.", "04-findings-and-human-review.png"),
        ("Agent activity trace", "Operational trace events show agent and tool activity without exposing hidden model reasoning.", "05-agent-activity-trace.png"),
    ]
    for title, caption, filename in screenshot_items:
        doc.add_heading(title, level=2)
        doc.add_paragraph(caption)
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(images[filename]), width=Inches(6.35))


def add_components(doc):
    doc.add_heading("Agent Components", level=1)
    doc.add_heading("Code Review Agent", level=2)
    doc.add_paragraph(
        "The Code Review Agent performs contextual analysis over bounded repository evidence. It can inspect Git diff, read constrained files, search the repository, inspect package metadata, find related tests, read coding standards, and consume deterministic quality results. It cannot execute arbitrary shell commands."
    )
    add_table(
        doc,
        ["Tool", "Purpose"],
        [
            ["getGitDiff", "Review changed code rather than the entire repository."],
            ["readFile", "Read bounded files under the configured repository root."],
            ["searchRepository", "Find related code by query without path traversal."],
            ["getPackageInfo", "Understand scripts and dependencies."],
            ["getRelatedTests", "Locate likely tests for a changed file."],
            ["getCodingStandards", "Provide local review rules and expected practices."],
            ["getQualityResults", "Use deterministic outputs as context without changing them."],
        ],
        widths=[1.65, 5.05],
    )
    doc.add_heading("Quality Agent", level=2)
    doc.add_paragraph(
        "The Quality Agent runs allowed deterministic commands and normalizes their outputs. It never invents results and treats tool infrastructure errors separately from quality failures."
    )
    add_table(
        doc,
        ["Tool", "Command", "Normalized output"],
        [
            ["ESLint", "npm run lint -- --format json", "Errors, warnings, file, line, rule, message"],
            ["TypeScript", "npm run typecheck", "Compilation status and compiler errors"],
            ["Jest and coverage", "npm run test:coverage -- --json", "Passed, failed, skipped tests and coverage percentages"],
            ["npm audit", "npm audit --json --audit-level=moderate", "Critical, high, moderate, low, and total vulnerabilities"],
        ],
        widths=[1.35, 2.45, 2.9],
    )


def add_api_and_models(doc):
    doc.add_heading("API and Data Model", level=1)
    add_table(
        doc,
        ["Endpoint", "Method", "Purpose"],
        [
            ["/api/health", "GET", "Health check"],
            ["/api/reviews", "POST", "Start a new review"],
            ["/api/reviews/:id", "GET", "Fetch complete review result"],
            ["/api/reviews/:id/findings", "GET", "Fetch review findings"],
            ["/api/reviews/:id/quality", "GET", "Fetch deterministic quality results"],
            ["/api/reviews/:id/trace", "GET", "Fetch operational trace"],
            ["/api/reviews/:reviewId/findings/:findingId", "PATCH", "Update finding status"],
        ],
        widths=[2.6, 0.8, 3.3],
    )
    doc.add_heading("Quality Check Model", level=2)
    doc.add_paragraph("Each deterministic check is normalized with a consistent shape.")
    add_table(
        doc,
        ["Field", "Meaning"],
        [
            ["tool", "Machine-readable tool id such as eslint, typecheck, jest, coverage, or npmAudit."],
            ["status", "PASS, FAIL, WARNING, ERROR, or SKIPPED."],
            ["durationMs", "Tool execution duration."],
            ["summary", "Human-readable result summary."],
            ["metrics", "Structured counts such as errors, warnings, failedTests, and vulnerabilities."],
            ["details", "Parsed per-file, per-test, or per-package details."],
        ],
        widths=[1.5, 5.2],
    )
    doc.add_heading("Policy Decisions", level=2)
    add_table(
        doc,
        ["Condition", "Decision"],
        [
            ["TypeScript/build failure", "BLOCKED"],
            ["Unit-test failure", "BLOCKED"],
            ["Critical or high deterministic security finding", "BLOCKED"],
            ["Configured deterministic quality gate failure", "BLOCKED"],
            ["AI critical or high finding", "HUMAN REVIEW REQUIRED"],
            ["AI medium finding", "WARNING"],
            ["All gates pass and no open high-risk AI findings remain", "PASS"],
        ],
        widths=[4.1, 2.6],
    )


def add_security_observability(doc):
    doc.add_heading("Security and Observability", level=1)
    doc.add_heading("Security Guardrails", level=2)
    add_bullet_list(
        doc,
        [
            "API keys are loaded from environment variables and .env is ignored by Git.",
            "Repository root validation prevents path traversal outside the workspace.",
            "Repository access for the AI agent is read-only.",
            "The LLM is not given arbitrary shell execution.",
            "Quality commands are selected from a fixed allowlist.",
            "AI output is validated with shared Zod schemas.",
            "Deterministic tool results remain authoritative.",
            "AI findings require human validation before release decisions.",
        ],
    )
    doc.add_heading("Trace Events", level=2)
    doc.add_paragraph(
        "The trace captures operational activity such as Review Started, Collecting Git Context, eslint, typecheck, jest, npmAudit, readFile, getRelatedTests, Policy Evaluation, and Review Completed. It intentionally excludes hidden model reasoning."
    )


def add_operations(doc):
    doc.add_heading("Setup and Operations", level=1)
    add_table(
        doc,
        ["Command", "Purpose"],
        [
            ["npm install", "Install workspace dependencies"],
            ["Copy-Item .env.example .env", "Create local environment file on Windows PowerShell"],
            ["npm run dev", "Run API and frontend together"],
            ["npm run build", "Build shared, agents, API, and web packages"],
            ["npm run test", "Run platform tests"],
            ["npm run lint", "Run platform linting"],
            ["npm run demo:quality", "Run deterministic quality checks against sample-project"],
        ],
        widths=[2.4, 4.3],
    )
    doc.add_heading("Environment Variables", level=2)
    add_table(
        doc,
        ["Variable", "Purpose", "Example"],
        [
            ["PORT", "API port", "3001"],
            ["WEB_ORIGIN", "CORS origin", "http://localhost:5173"],
            ["REPOSITORY_ROOT", "Repository to review", "./sample-project"],
            ["GEMINI_API_KEY", "Gemini API key", "your_key_here"],
            ["GEMINI_MODEL", "Gemini model", "gemini-3.5-flash-lite"],
            ["ENABLE_LLM_REVIEW", "Enable live LLM calls", "true"],
            ["MAX_REVIEW_FILES", "Changed file cap", "30"],
            ["MAX_FILE_BYTES", "File read cap", "12000"],
        ],
        widths=[1.65, 3.2, 1.85],
    )


def add_ground_truth(doc):
    doc.add_heading("Sample Project Ground Truth", level=1)
    doc.add_paragraph(
        "The sample project intentionally contains defects that demonstrate the value of combining deterministic tools with contextual AI review."
    )
    add_table(
        doc,
        ["ID", "Intentional defect", "Expected detector"],
        [
            ["GT-01", "userId trusted from request body instead of authenticated principal", "Code Review Agent"],
            ["GT-02", "DTO missing validation decorators and validation pipe", "Code Review Agent"],
            ["GT-03", "Zero quantity allowed", "Code Review Agent"],
            ["GT-04", "Restricted product authorization missing", "Code Review Agent"],
            ["GT-05", "Sensitive-looking context logged with console.log", "ESLint and Code Review Agent"],
            ["GT-06", "Exceptions swallowed in findOrdersForUser", "Code Review Agent and ESLint"],
            ["GT-07", "Sequential awaits in summary logic", "Code Review Agent"],
            ["GT-08", "Coupon logic returns zero for no-discount orders", "Jest and Code Review Agent"],
            ["GT-09", "Missing edge-case tests", "Code Review Agent"],
            ["GT-10", "Old demo dependency tree with real audit findings", "npm audit"],
        ],
        widths=[0.7, 4.1, 1.9],
    )


def add_testing_and_acceptance(doc):
    doc.add_heading("Testing and Acceptance", level=1)
    doc.add_paragraph(
        "The platform verification pass ran build, test, lint, deterministic quality demo, and end-to-end API review execution."
    )
    add_table(
        doc,
        ["Verification", "Result"],
        [
            ["npm run build", "Passed"],
            ["npm run test", "Passed"],
            ["npm run lint", "Passed"],
            ["npm run demo:quality", "Executed real sample-project quality checks"],
            ["End-to-end API review", "Completed with BLOCKED decision and real quality results"],
            ["Human finding update", "PATCH endpoint verified and policy recalculated"],
        ],
        widths=[2.4, 4.3],
    )
    doc.add_heading("Acceptance Criteria Summary", level=2)
    add_table(
        doc,
        ["Capability", "Status"],
        [
            ["React UI runs", "Complete"],
            ["NestJS backend runs", "Complete"],
            ["User can start review", "Complete"],
            ["Code Review Agent executes", "Complete"],
            ["Quality Agent executes", "Complete"],
            ["Real ESLint, TypeScript, Jest, coverage, and audit results collected", "Complete"],
            ["Policy engine produces final decision", "Complete"],
            ["UI displays findings, quality checks, and trace", "Complete"],
            ["Human can accept and dismiss findings", "Complete"],
            ["Secrets excluded from Git", "Complete"],
        ],
        widths=[4.7, 2.0],
    )


def add_limitations_and_roadmap(doc):
    doc.add_heading("Limitations and Roadmap", level=1)
    add_table(
        doc,
        ["Current limitation", "Recommended future improvement"],
        [
            ["Review state is in-memory", "Persist reviews and audit events in MongoDB or Postgres"],
            ["Endpoint is synchronous", "Move review execution to queue-backed workers"],
            ["No GitHub PR integration", "Add GitHub Actions, webhook, PR comments, and status checks"],
            ["No authentication", "Add SSO and role-based reviewer permissions"],
            ["Single-language demo focus", "Add language-specific quality runners and policy profiles"],
            ["No auto-fix flow", "Generate patches only after explicit human approval"],
            ["Basic cost controls", "Add model routing, token budgets, caching, and review summaries"],
        ],
        widths=[3.1, 3.6],
    )


def prepare_images():
    vertical_offsets = {
        "01-dashboard-ready.png": 0,
        "02-dashboard-review-complete.png": 0,
        "03-quality-checks.png": 390,
        "04-findings-and-human-review.png": 520,
        "05-agent-activity-trace.png": 760,
    }
    images = {}
    for image_path in SCREEN_DIR.glob("0*.png"):
        with Image.open(image_path) as img:
            img = img.convert("RGB")
            target_ratio = 16 / 10
            current_ratio = img.width / img.height
            if current_ratio < target_ratio:
                crop_height = min(img.height, int(img.width / target_ratio))
                y_offset = min(vertical_offsets.get(image_path.name, 0), max(0, img.height - crop_height))
                img = img.crop((0, y_offset, img.width, y_offset + crop_height))
            elif current_ratio > target_ratio:
                crop_width = min(img.width, int(img.height * target_ratio))
                img = img.crop((0, 0, crop_width, img.height))
            img = ImageOps.contain(img, (1800, 1100), method=Image.Resampling.LANCZOS)
            canvas = Image.new("RGB", (1800, 1100), "white")
            x = (1800 - img.width) // 2
            y = (1100 - img.height) // 2
            canvas.paste(img, (x, y))
            output = DOCX_IMG_DIR / image_path.name
            canvas.save(output, "PNG", optimize=True)
            images[image_path.name] = output
    return images


def add_key_value_table(doc, rows):
    add_table(doc, ["Item", "Details"], rows, widths=[1.8, 4.9])


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, "1F2937")
        set_cell_text(cell, header, bold=True, color="FFFFFF", center=index == 0 and len(headers) > 2)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            set_cell_text(cell, str(value), center=index == 0 and len(headers) > 2)
    set_widths(table, widths)
    doc.add_paragraph()
    return table


def set_cell_text(cell, text, bold=False, color="000000", center=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(color)
    set_cell_margins(cell)


def set_widths(table, widths):
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)


def set_cell_margins(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), "90")
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "D9D9D9")
        borders.append(tag)
    tbl_pr.append(borders)


def add_bullet_list(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered_list(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


if __name__ == "__main__":
    main()
